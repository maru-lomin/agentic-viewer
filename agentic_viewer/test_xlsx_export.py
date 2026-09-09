"""Tests for evaluation XLSX export."""

from __future__ import annotations

import io
import json
import tempfile
import unittest
from pathlib import Path

import openpyxl
from agentic_viewer.evaluation.xlsx_export import (
    EXCEL_COLUMNS,
    extract_run_rows,
    generate_evaluation_xlsx,
)
from agentic_viewer.routers.evaluations import (
    get_export_evaluation_xlsx,
    post_export_evaluation_xlsx,
)
from fastapi import HTTPException


class XlsxExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.runs_root = Path(self.tmp.name) / "runs"
        self.runs_root.mkdir(parents=True)

        # Create sample run 1
        self.run1_dir = self.runs_root / "run-sample-1"
        self.run1_dir.mkdir()
        (self.run1_dir / "meta.json").write_text(
            json.dumps({
                "run_id": "run-sample-1",
                "source_filename": "sample_doc_1.pdf",
                "status": "ok",
            }),
            encoding="utf-8",
        )
        (self.run1_dir / "04_result.json").write_text(
            json.dumps({
                "kv_results": [
                    {
                        "key": "Distance between GSU Transformers",
                        "value": "20m",
                        "search_pages": [11],
                        "evidence_quote": "11페이지에 20m로 기재됨",
                        "search_reasons": {"11": "11페이지에서 트랜스포머 이격거리 발견"},
                    },
                    {
                        "key": "Configuration",
                        "value": "2x1 Combined Cycle",
                        "search_pages": [4, 5],
                        "value_reason": "4페이지 개요에서 확인",
                        "search_reasons": {"4": "4페이지 개요", "5": "5페이지 상세"},
                    },
                ]
            }),
            encoding="utf-8",
        )
        (self.run1_dir / "05_eval.json").write_text(
            json.dumps({
                "document": "sample_doc_1.pdf",
                "per_key": [
                    {
                        "key": "Distance between GSU Transformers",
                        "value": {"pred": "20m", "gold": "15~20m", "exact_match": False},
                    },
                    {
                        "key": "Configuration",
                        "value": {"pred": "2x1 Combined Cycle", "gold": "2x1 Combined Cycle", "exact_match": True},
                    },
                ]
            }),
            encoding="utf-8",
        )
        agentic_dir1 = self.run1_dir / "06_agentic_eval"
        agentic_dir1.mkdir()
        (agentic_dir1 / "Distance_between_GSU_Transformers.json").write_text(
            json.dumps({
                "key": "Distance between GSU Transformers",
                "status": "done",
                "is_correct_answer": "incorrect",
                "reason_summary": "골드값과 범위 불일치",
                "reason_detail": "문서 상세 근거에 따르면...",
            }),
            encoding="utf-8",
        )
        (agentic_dir1 / "Configuration.json").write_text(
            json.dumps({
                "key": "Configuration",
                "status": "done",
                "is_correct_answer": "correct",
                "reason_summary": "정답 일치",
                "reason_detail": "개요와 일치함",
            }),
            encoding="utf-8",
        )

        # Create sample run 2
        self.run2_dir = self.runs_root / "run-sample-2"
        self.run2_dir.mkdir()
        (self.run2_dir / "meta.json").write_text(
            json.dumps({
                "run_id": "run-sample-2",
                "source_filename": "sample_doc_2.pdf",
                "status": "ok",
            }),
            encoding="utf-8",
        )
        (self.run2_dir / "04_result.json").write_text(
            json.dumps({
                "kv_results": [
                    {
                        "key": "LTSA",
                        "value": "Yes",
                        "search_pages": [18],
                        "evidence_quote": "18페이지 LTSA 체결 명시",
                        "search_reasons": {"18": "LTSA 섹션"},
                    }
                ]
            }),
            encoding="utf-8",
        )

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_extract_run_rows(self) -> None:
        rows = extract_run_rows(self.run1_dir)
        self.assertEqual(len(rows), 2)

        r1 = rows[0]
        self.assertEqual(r1["파일명"], "sample_doc_1.pdf")
        self.assertEqual(r1["Key"], "Distance between GSU Transformers")
        self.assertEqual(r1["Prediction"], "20m")
        self.assertEqual(r1["Ground Truth"], "15~20m")
        self.assertEqual(r1["채점결과"], "오답")
        self.assertEqual(r1["검색페이지"], "11")
        self.assertEqual(r1["추출근거"], "11페이지에 20m로 기재됨")
        self.assertIn("p.11", r1["검색근거"])
        self.assertEqual(r1["검토의견-요약"], "골드값과 범위 불일치")
        self.assertEqual(r1["검토의견-상세"], "문서 상세 근거에 따르면...")

        r2 = rows[1]
        self.assertEqual(r2["Key"], "Configuration")
        self.assertEqual(r2["채점결과"], "정답")
        self.assertEqual(r2["검색페이지"], "4, 5")

    def test_generate_evaluation_xlsx(self) -> None:
        buf = generate_evaluation_xlsx([self.run1_dir, self.run2_dir])
        data = buf.getvalue()
        self.assertGreater(len(data), 500)

        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb.active
        self.assertEqual(ws.title, "채점결과")

        headers = [ws.cell(row=1, column=c).value for c in range(1, len(EXCEL_COLUMNS) + 1)]
        self.assertEqual(headers, EXCEL_COLUMNS)

        # 3 data rows (2 from run1, 1 from run2) + 1 header row = 4 rows
        self.assertEqual(ws.max_row, 4)
        self.assertEqual(ws.cell(row=2, column=1).value, "sample_doc_1.pdf")
        self.assertEqual(ws.cell(row=4, column=1).value, "sample_doc_2.pdf")

    def test_api_export_xlsx(self) -> None:
        import agentic_viewer.app as app_mod
        orig_runs_root = app_mod.RUNS_ROOT
        app_mod.RUNS_ROOT = self.runs_root
        try:
            # Test GET
            resp = get_export_evaluation_xlsx(run_ids="run-sample-1,run-sample-2")
            self.assertEqual(
                resp.media_type,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            self.assertIn("attachment", resp.headers["Content-Disposition"])
            self.assertIn(".xlsx", resp.headers["Content-Disposition"])

            wb = openpyxl.load_workbook(io.BytesIO(resp.body))
            self.assertEqual(wb.active.max_row, 4)

            # Test POST
            resp_post = post_export_evaluation_xlsx(body={"run_ids": ["run-sample-1"]})
            self.assertEqual(
                resp_post.media_type,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            wb_post = openpyxl.load_workbook(io.BytesIO(resp_post.body))
            self.assertEqual(wb_post.active.max_row, 3)

            # Test 400 on empty
            with self.assertRaises(HTTPException) as cm:
                get_export_evaluation_xlsx(run_ids="")
            self.assertEqual(cm.exception.status_code, 400)

            # Test 404 on nonexistent
            with self.assertRaises(HTTPException) as cm:
                get_export_evaluation_xlsx(run_ids="nonexistent-run")
            self.assertEqual(cm.exception.status_code, 404)
        finally:
            app_mod.RUNS_ROOT = orig_runs_root


if __name__ == "__main__":
    unittest.main()
