"""Export evaluation and inference results to Excel (.xlsx) format."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from agentic_viewer.evaluation.baseline import load_or_compute_run_eval
from agentic_viewer.evaluation.summary import read_agentic_evals
from agentic_viewer.pdf_source import infer_run_document

EXCEL_COLUMNS = [
    "파일명",
    "Key",
    "Prediction",
    "Ground Truth",
    "채점결과",
    "검색페이지",
    "추출근거",
    "검색근거",
    "검토의견-요약",
    "검토의견-상세",
]


def _read_json(path: Path) -> Any:
    if not path.is_file():
        return None
    try:
        import json
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None


def _format_search_pages(pages_val: Any) -> str:
    if not pages_val:
        return ""
    if isinstance(pages_val, (list, tuple, set)):
        sorted_pages = []
        for p in pages_val:
            try:
                sorted_pages.append((int(p), str(p)))
            except (ValueError, TypeError):
                sorted_pages.append((999999, str(p)))
        sorted_pages.sort(key=lambda x: x[0])
        return ", ".join(x[1] for x in sorted_pages)
    return str(pages_val)


def _format_search_reasons(reasons_val: Any) -> str:
    if not reasons_val:
        return ""
    if isinstance(reasons_val, dict):
        lines = []
        items = []
        for k, v in reasons_val.items():
            try:
                page_num = int(k)
            except (ValueError, TypeError):
                page_num = 999999
            items.append((page_num, str(k), str(v or "").strip()))
        items.sort(key=lambda x: x[0])
        for _, k_str, v_str in items:
            if v_str:
                lines.append(f"p.{k_str}: {v_str}")
        return "\n".join(lines)
    return str(reasons_val).strip()


def _format_verdict(is_correct: Optional[str], exact_match: Optional[bool], status: Optional[str]) -> str:
    if is_correct:
        norm = str(is_correct).strip().lower()
        if norm == "correct":
            return "정답"
        if norm == "incorrect":
            return "오답"
        return str(is_correct).strip()

    if status == "running":
        return "평가중"

    if exact_match is True:
        return "정답 (EM)"
    if exact_match is False:
        return "오답 (EM)"

    return "미채점"


def extract_run_rows(run_dir: Path) -> List[Dict[str, str]]:
    """Extract evaluation and extraction rows for a single run directory."""
    if not run_dir.is_dir():
        return []

    doc_name = infer_run_document(run_dir) or run_dir.name
    res_data = _read_json(run_dir / "04_result.json") or {}
    kv_results = res_data.get("kv_results") or []

    pred_by_key: Dict[str, Dict[str, Any]] = {}
    if isinstance(kv_results, list):
        for item in kv_results:
            if isinstance(item, dict) and item.get("key"):
                pred_by_key[str(item["key"])] = item

    eval_data = _read_json(run_dir / "05_eval.json")
    if not isinstance(eval_data, dict) or not eval_data.get("per_key"):
        computed = load_or_compute_run_eval(run_dir)
        if isinstance(computed, dict) and computed.get("per_key"):
            eval_data = computed
        elif not isinstance(eval_data, dict):
            eval_data = {}

    per_key = eval_data.get("per_key") or []
    eval_by_key: Dict[str, Dict[str, Any]] = {}
    if isinstance(per_key, list):
        for item in per_key:
            if isinstance(item, dict) and item.get("key"):
                eval_by_key[str(item["key"])] = item

    agentic_by_key = read_agentic_evals(run_dir)

    # Collect keys preserving original order
    ordered_keys: List[str] = []
    seen_keys = set()

    for item in kv_results:
        if isinstance(item, dict) and item.get("key"):
            k = str(item["key"])
            if k not in seen_keys:
                seen_keys.add(k)
                ordered_keys.append(k)

    for item in per_key:
        if isinstance(item, dict) and item.get("key"):
            k = str(item["key"])
            if k not in seen_keys:
                seen_keys.add(k)
                ordered_keys.append(k)

    for k in agentic_by_key:
        k_str = str(k)
        if k_str not in seen_keys:
            seen_keys.add(k_str)
            ordered_keys.append(k_str)

    rows: List[Dict[str, str]] = []
    for key in ordered_keys:
        p_item = pred_by_key.get(key, {})
        e_item = eval_by_key.get(key, {})
        a_item = agentic_by_key.get(key, {})

        # 1. Prediction
        pred = p_item.get("value")
        if pred is None:
            pred = e_item.get("value", {}).get("pred")
        pred_str = str(pred) if pred is not None else ""

        # 2. Ground Truth
        gt = e_item.get("value", {}).get("gold")
        gt_str = str(gt) if gt is not None else ""

        # 3. 채점결과
        is_correct = a_item.get("is_correct_answer")
        exact_match = e_item.get("value", {}).get("exact_match")
        status = a_item.get("status")
        verdict = _format_verdict(is_correct, exact_match, status)

        # 4. 검색페이지
        search_pages = p_item.get("search_pages")
        if not search_pages and isinstance(e_item.get("search_pages"), dict):
            search_pages = e_item["search_pages"].get("pred")
        pages_str = _format_search_pages(search_pages)

        # 5. 추출근거
        evidence = p_item.get("evidence_quote") or p_item.get("value_reason")
        if not evidence:
            raw_ev = p_item.get("evidence")
            if isinstance(raw_ev, list):
                evidence = "\n".join(
                    str(x.get("text", "")).strip()
                    for x in raw_ev
                    if isinstance(x, dict) and x.get("text")
                )
        evidence_str = str(evidence).strip() if evidence else ""

        # 6. 검색근거
        search_reasons = p_item.get("search_reasons")
        reasons_str = _format_search_reasons(search_reasons)

        # 7. 검토의견-요약
        reason_summary = a_item.get("reason_summary") or a_item.get("reason") or ""
        reason_summary_str = str(reason_summary).strip()

        # 8. 검토의견-상세
        reason_detail = a_item.get("reason_detail") or a_item.get("text") or ""
        reason_detail_str = str(reason_detail).strip()

        rows.append({
            "파일명": doc_name,
            "Key": key,
            "Prediction": pred_str,
            "Ground Truth": gt_str,
            "채점결과": verdict,
            "검색페이지": pages_str,
            "추출근거": evidence_str,
            "검색근거": reasons_str,
            "검토의견-요약": reason_summary_str,
            "검토의견-상세": reason_detail_str,
        })

    return rows


def generate_evaluation_xlsx(run_dirs: Sequence[Path]) -> io.BytesIO:
    """Generate Excel workbook from multiple run directories."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "채점결과"

    # Header styling
    header_font = Font(name="Malgun Gothic", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border_side = Side(style="thin", color="D3D3D3")
    data_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )
    header_border = Border(
        left=Side(style="thin", color="0D233A"),
        right=Side(style="thin", color="0D233A"),
        top=Side(style="thin", color="0D233A"),
        bottom=Side(style="medium", color="0D233A"),
    )

    # Verdict highlight fills
    correct_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # soft green
    incorrect_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")  # soft red/orange
    evaluating_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # soft yellow

    # Write headers
    ws.append(EXCEL_COLUMNS)
    ws.row_dimensions[1].height = 28

    for col_idx in range(1, len(EXCEL_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = header_border

    # Alignment presets by column
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align_mid = Alignment(horizontal="left", vertical="center", wrap_text=True)
    left_align_top = Alignment(horizontal="left", vertical="top", wrap_text=True)
    data_font = Font(name="Malgun Gothic", size=10)

    row_idx = 2
    for rdir in run_dirs:
        rows = extract_run_rows(rdir)
        for rdata in rows:
            row_values = [rdata.get(col, "") for col in EXCEL_COLUMNS]
            ws.append(row_values)

            # Style each cell in row
            verdict = rdata.get("채점결과", "")
            for col_idx, col_name in enumerate(EXCEL_COLUMNS, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = data_border

                if col_name in ("파일명", "Key", "채점결과", "검색페이지"):
                    cell.alignment = center_align
                elif col_name in ("Prediction", "Ground Truth"):
                    cell.alignment = left_align_mid
                else:  # 추출근거, 검색근거, 검토의견-요약, 검토의견-상세
                    cell.alignment = left_align_top

                # Verdict cell coloring
                if col_name == "채점결과":
                    if verdict in ("정답", "정답 (EM)"):
                        cell.fill = correct_fill
                        cell.font = Font(name="Malgun Gothic", size=10, bold=True, color="276A3C")
                    elif verdict in ("오답", "오답 (EM)"):
                        cell.fill = incorrect_fill
                        cell.font = Font(name="Malgun Gothic", size=10, bold=True, color="C00000")
                    elif verdict == "평가중":
                        cell.fill = evaluating_fill
                        cell.font = Font(name="Malgun Gothic", size=10, bold=True, color="B25E00")

            row_idx += 1

    # Freeze header pane
    ws.freeze_panes = "A2"

    # Set column widths
    column_widths = {
        "파일명": 32,
        "Key": 24,
        "Prediction": 20,
        "Ground Truth": 20,
        "채점결과": 12,
        "검색페이지": 12,
        "추출근거": 42,
        "검색근거": 42,
        "검토의견-요약": 36,
        "검토의견-상세": 55,
    }
    for col_idx, col_name in enumerate(EXCEL_COLUMNS, start=1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = column_widths.get(col_name, 20)

    # Enable auto filter
    last_col_letter = get_column_letter(len(EXCEL_COLUMNS))
    total_rows = max(row_idx - 1, 1)
    ws.auto_filter.ref = f"A1:{last_col_letter}{total_rows}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
